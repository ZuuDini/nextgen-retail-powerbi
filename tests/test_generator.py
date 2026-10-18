"""Generator tests: determinism, every planted defect present, cleanup recovers truth."""

from __future__ import annotations

import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from generator import config as C
from generator.expected_kpis import compute
from generator.generate import run
from generator.reference_cleanup import clean_all


@pytest.fixture(scope="session")
def gen(tmp_path_factory):
    out = tmp_path_factory.mktemp("run1")
    return out, run(C.DEFAULT_SEED, out)


@pytest.fixture(scope="session")
def raw(gen):
    return gen[0] / "raw"


@pytest.fixture(scope="session")
def cleaned(raw):
    return clean_all(raw)


def _xlsx_cells(path):
    wb = load_workbook(path)
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb}


# ---------- determinism ----------------------------------------------------

def test_same_seed_same_files(gen, tmp_path):
    out1, _ = gen
    run(C.DEFAULT_SEED, tmp_path)
    for name in ("sales_transactions_2024.csv", "sales_transactions_2025.csv",
                 "inventory_snapshots.csv", "customers.json"):
        assert (out1 / "raw" / name).read_bytes() == (tmp_path / "raw" / name).read_bytes(), name
    for name in ("product_master.xlsx", "regional_performance.xlsx"):
        assert _xlsx_cells(out1 / "raw" / name) == _xlsx_cells(tmp_path / "raw" / name), name
    assert (out1 / "expected" / "kpis.json").read_text() == \
        (tmp_path / "expected" / "kpis.json").read_text()


def test_committed_answer_key_matches_default_seed(gen):
    """data/expected/kpis.json in the repo must be the seed-387 output."""
    from generator.generate import ROOT
    committed = json.loads((ROOT / "data" / "expected" / "kpis.json").read_text())
    assert committed == gen[1]["kpis"]


# ---------- defects are actually present -----------------------------------

def test_sales_defects(raw):
    s24 = pd.read_csv(raw / "sales_transactions_2024.csv", dtype=str, keep_default_na=False)
    s25 = pd.read_csv(raw / "sales_transactions_2025.csv", dtype=str, keep_default_na=False)
    assert s24["OrderID"].iloc[-1] == "TOTAL"
    assert s24["OrderDate"].str.match(r"^\d{2}/\d{2}/\d{4}$").iloc[:-1].all()
    assert s25["OrderDate"].str.match(r"^\d{4}-\d{2}-\d{2}$").any()
    assert s25["OrderDate"].str.match(r"^\d{2}-[A-Z][a-z]{2}-\d{4}$").any()
    assert s25["UnitPrice"].str.startswith("$").all()
    assert s25["UnitPrice"].str.contains(",").any()
    assert s25["DiscountPct"].str.endswith("%").all()
    assert s24["DiscountPct"].iloc[:-1].str.match(r"^0\.\d{2}$").all()
    assert s25.duplicated().sum() > 100
    assert (s24["StoreCode"] != s24["StoreCode"].str.strip()).sum() > 100
    assert (s24["Quantity"].iloc[:-1].astype(int) < 0).any()      # returns
    assert ((s24["ShipDate"] == "") & (s24["OrderID"] != "TOTAL")).any()


def test_product_master_defects(raw):
    wb = load_workbook(raw / "product_master.xlsx")
    ws = wb["Products"]
    assert ws["A4"].value == "Product ID"
    assert isinstance(ws["A5"].value, int)                         # zeros lost
    assert len(ws.merged_cells.ranges) == 4
    assert any(str(ws.cell(r, 2).value).endswith(" ") for r in range(5, ws.max_row + 1))


def test_regional_defects(raw):
    cells = _xlsx_cells(raw / "regional_performance.xlsx")
    regions = {row[4] for row in cells["Stores"][1:]}
    assert len(regions) > 4
    targets = cells["Monthly Targets"]
    assert targets[2][0] == "Region" and len(targets[2]) == 25     # wide layout
    assert targets[-1][0] == "Total"


def test_customer_and_inventory_defects(raw, gen):
    doc = json.loads((raw / "customers.json").read_text())
    custs = doc["customers"]
    assert isinstance(custs[0]["demographics"], dict)
    assert sum(c["segment"] is None for c in custs) == gen[1]["profile"]["customers"]["null_segment"]
    assert {c["demographics"]["gender"] for c in custs} >= {"F", "Female", "female", "M"}
    inv = pd.read_csv(raw / "inventory_snapshots.csv", dtype=str, keep_default_na=False)
    assert inv["store"].str.match(r"^st\d{3}$").all()
    assert (inv["store"] == "st099").sum() == 720
    assert (inv["on_hand"] == "").sum() == gen[1]["profile"]["inventory"]["blank_on_hand"]


# ---------- the documented cleanup recovers the answer key -----------------

@pytest.mark.parametrize("table", ["stores", "products", "customers", "sales",
                                   "inventory", "targets"])
def test_reference_cleanup_matches_clean(cleaned, gen, table):
    got = cleaned[table].reset_index(drop=True)
    expected = gen[1]["clean"][table].reset_index(drop=True)
    pd.testing.assert_frame_equal(got, expected[got.columns], check_dtype=False)


def test_kpis_from_cleanup_equal_answer_key(cleaned, gen):
    assert compute(cleaned) == gen[1]["kpis"]


def test_profile_counts_add_up(gen):
    p, k = gen[1]["profile"], gen[1]["kpis"]
    s = p["sales"]
    clean = sum(s[y]["raw_rows"] - s[y]["duplicates"] - s[y]["footer_rows"] for y in s)
    assert clean == k["Row Counts"]["FactSales"]
    i = p["inventory"]
    assert i["raw_rows"] - i["blank_on_hand"] - i["orphan_store_rows"] == k["Row Counts"]["FactInventory"]


def test_planted_story_is_visible(gen):
    k = gen[1]["kpis"]["By Region and Year"]
    south, others = k["South"], [k[r] for r in ("Northeast", "Midwest", "West")]
    assert south["2025"]["Sales YoY %"] > 0.2
    assert south["2025"]["Margin %"] < south["2024"]["Margin %"] - 0.05
    assert south["2025"]["Inventory Turnover"] < min(o["2025"]["Inventory Turnover"] for o in others)
    seg = gen[1]["kpis"]["By Segment and Year"]["Loyalist"]
    assert seg["2025"]["Repeat Rate %"] < seg["2024"]["Repeat Rate %"]
